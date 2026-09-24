"""Downloadable PPTX version of the in-app "Ayuda" (help) content — the
"PPTX" button next to "Descargar PDF" in the Ayuda modal.

Same source of truth as help_report.py (the PDF export): the Ayuda text
lives entirely client-side as HTML template literals in
static/js/prionvault.js, so the frontend sends the already-rendered HTML
per tab and this module turns it into a designed slide deck — meaning
BOTH exports always reflect whatever the Ayuda content currently says,
with nothing to keep in sync by hand. Reuses help_report's tiny HTML
parser (_parse_html/_Node/_BLOCK_TAGS) instead of duplicating it.

Layout: one title slide, then per Ayuda tab a coloured divider slide
followed by one content slide per <h3> heading in that tab (splitting
into "(cont.)" slides if a section's content would overflow one slide),
plus a dedicated slide for any <table>.
"""
from __future__ import annotations

from .help_report import _parse_html, _Node, _BLOCK_TAGS, _has_block_children

ACCENT_HEX = "0F3460"
ACCENT_SOFT_HEX = "A5B4FC"
TEXT_HEX = "1F2937"
MUTED_HEX = "6B7280"
WHITE_HEX = "FFFFFF"

# Rough capacity of a content slide's body — one <li>/short <p> costs 1-2
# "units"; once a section's running total would exceed this, a new
# "(cont.)" slide starts instead of overflowing the slide.
_MAX_BODY_UNITS = 13


def _hex(h: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h)


def _rect(slide, prs, left, top, width, height, fill_hex):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _hex(fill_hex)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_footer(slide, prs, label: str = ""):
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(0.55), prs.slide_height - Inches(0.42),
                                   prs.slide_width - Inches(1.1), Inches(0.32))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"PrionVault · Guía de usuario{'  ·  ' + label if label else ''}"
    p.font.size = Pt(9)
    p.font.color.rgb = _hex(MUTED_HEX)


def _title_slide(prs, generated_at: str):
    from pptx.util import Inches, Pt
    slide = _blank_slide(prs)
    _rect(slide, prs, 0, 0, prs.slide_width, prs.slide_height, ACCENT_HEX)

    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), prs.slide_width - Inches(2), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔬 PrionVault"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = _hex(WHITE_HEX)
    p2 = tf.add_paragraph()
    p2.text = "Guía completa de usuario"
    p2.font.size = Pt(26)
    p2.font.color.rgb = _hex(ACCENT_SOFT_HEX)

    sub_box = slide.shapes.add_textbox(Inches(1), Inches(6.5), prs.slide_width - Inches(2), Inches(0.5))
    tf2 = sub_box.text_frame
    p3 = tf2.paragraphs[0]
    p3.text = f"Generado automáticamente desde la Ayuda de la app el {generated_at}"
    p3.font.size = Pt(13)
    p3.font.color.rgb = _hex(ACCENT_SOFT_HEX)
    return slide


def _divider_slide(prs, label: str):
    from pptx.util import Inches, Pt
    slide = _blank_slide(prs)
    _rect(slide, prs, 0, 0, prs.slide_width, prs.slide_height, ACCENT_HEX)
    box = slide.shapes.add_textbox(Inches(1), Inches(3.15), prs.slide_width - Inches(2), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = _hex(WHITE_HEX)
    return slide


def _content_slide(prs, title: str):
    from pptx.util import Inches, Pt
    slide = _blank_slide(prs)
    _rect(slide, prs, 0, 0, prs.slide_width, Inches(0.12), ACCENT_HEX)

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), prs.slide_width - Inches(1.2), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = _hex(ACCENT_HEX)

    body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.35),
                                        prs.slide_width - Inches(1.4),
                                        prs.slide_height - Inches(1.9))
    body_tf = body_box.text_frame
    body_tf.word_wrap = True
    return slide, body_tf


def _apply_runs(paragraph, runs, size_pt, color_hex, bold_default=False):
    from pptx.util import Pt
    if not runs:
        run = paragraph.add_run()
        run.text = ""
        return
    for text, fmt in runs:
        if not text:
            continue
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(size_pt)
        run.font.bold = bold_default or bool(fmt.get("bold"))
        run.font.italic = bool(fmt.get("italic"))
        run.font.color.rgb = _hex(color_hex)
        if fmt.get("code"):
            run.font.name = "Consolas"


def _inline_runs(node) -> list[tuple[str, dict]]:
    """Like help_report._inline_markup but returns (text, format) runs for
    pptx instead of ReportLab markup text."""
    if isinstance(node, str):
        return [(node, {})] if node.strip() else []
    out: list[tuple[str, dict]] = []
    for child in node.children:
        if isinstance(child, str):
            if child:
                out.append((child, {}))
            continue
        tag = child.tag
        if tag in ("strong", "b"):
            out.extend((t, {**f, "bold": True}) for t, f in _inline_runs(child))
        elif tag == "em":
            out.extend((t, {**f, "italic": True}) for t, f in _inline_runs(child))
        elif tag == "code":
            out.extend((t, {**f, "code": True}) for t, f in _inline_runs(child))
        elif tag == "br":
            out.append((" ", {}))
        elif tag in _BLOCK_TAGS:
            continue
        else:  # span, i, a, ... — transparent inline containers
            out.extend(_inline_runs(child))
    return out


def _runs_text(runs: list[tuple[str, dict]]) -> str:
    return "".join(t for t, _ in runs).strip()


class _SlideWriter:
    """Accumulates one Ayuda section's content across as many slides as
    needed, starting a "(cont.)" slide whenever the running body-unit
    total would overflow the current one."""

    def __init__(self, prs, title: str, section_label: str):
        self.prs = prs
        self.base_title = title or section_label
        self.section_label = section_label
        self.units = 0
        self.part = 1
        self.slide, self.tf = _content_slide(prs, self._title_text())
        _add_footer(self.slide, prs, section_label)
        self._first_para = True

    def _title_text(self) -> str:
        return self.base_title if self.part == 1 else f"{self.base_title} (cont.)"

    def _ensure_room(self, cost: int):
        if self.units > 0 and self.units + cost > _MAX_BODY_UNITS:
            self.part += 1
            self.slide, self.tf = _content_slide(self.prs, self._title_text())
            _add_footer(self.slide, self.prs, self.section_label)
            self.units = 0
            self._first_para = True

    def _next_paragraph(self):
        if self._first_para:
            self._first_para = False
            return self.tf.paragraphs[0]
        return self.tf.add_paragraph()

    def add_h4(self, runs):
        if not runs:
            return
        from pptx.util import Pt
        self._ensure_room(2)
        p = self._next_paragraph()
        p.space_before = Pt(10)
        _apply_runs(p, runs, 15, ACCENT_HEX, bold_default=True)
        self.units += 2

    def add_p(self, runs):
        if not runs:
            return
        from pptx.util import Pt
        self._ensure_room(2)
        p = self._next_paragraph()
        p.space_before = Pt(5)
        _apply_runs(p, runs, 13, TEXT_HEX)
        self.units += 2

    def add_li(self, runs, level: int = 0):
        if not runs:
            return
        from pptx.util import Pt
        self._ensure_room(1)
        p = self._next_paragraph()
        p.level = min(level, 4)
        p.space_before = Pt(2)
        _apply_runs(p, [("•  ", {})] + runs, 13, TEXT_HEX)
        self.units += 1


def _table_slide(prs, title: str, section_label: str, table_node: _Node):
    """Native pptx table on its own dedicated slide (colspan is ignored —
    the Ayuda tables in practice are simple grids)."""
    from pptx.util import Inches, Pt

    rows: list[_Node] = []

    def walk(node: _Node):
        for c in node.children:
            if isinstance(c, _Node):
                if c.tag == "tr":
                    rows.append(c)
                elif c.tag in ("thead", "tbody"):
                    walk(c)

    walk(table_node)
    if not rows:
        return

    grid: list[list[str]] = []
    header_row_idx = None
    for ri, tr in enumerate(rows):
        cells = [c for c in tr.children if isinstance(c, _Node) and c.tag in ("td", "th")]
        if not cells:
            continue
        if header_row_idx is None and all(c.tag == "th" for c in cells):
            header_row_idx = len(grid)
        grid.append([_runs_text(_inline_runs(c)) for c in cells])
    if not grid:
        return
    max_cols = max(len(r) for r in grid)
    for r in grid:
        while len(r) < max_cols:
            r.append("")

    slide, _unused = _content_slide(prs, title)
    _add_footer(slide, prs, section_label)

    n_rows, n_cols = len(grid), max_cols
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.7), Inches(1.4),
        prs.slide_width - Inches(1.4), Inches(5.4),
    ).table
    for ri, row in enumerate(grid):
        for ci in range(n_cols):
            cell = table_shape.cell(ri, ci)
            cell.text = row[ci] if ci < len(row) else ""
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(11)
                if ri == header_row_idx:
                    para.font.bold = True
                    para.font.color.rgb = _hex(WHITE_HEX)
                else:
                    para.font.color.rgb = _hex(TEXT_HEX)
            if ri == header_row_idx:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _hex(ACCENT_HEX)


def _render_section(prs, root: _Node, section_label: str):
    state = {"writer": None}

    def writer() -> _SlideWriter:
        if state["writer"] is None:
            state["writer"] = _SlideWriter(prs, section_label, section_label)
        return state["writer"]

    def walk(node: _Node):
        for child in node.children:
            if isinstance(child, str):
                txt = child.strip()
                if txt:
                    writer().add_p([(txt, {})])
                continue

            tag = child.tag
            if tag == "h3":
                title_text = _runs_text(_inline_runs(child)) or section_label
                state["writer"] = _SlideWriter(prs, title_text, section_label)
            elif tag == "h4":
                writer().add_h4(_inline_runs(child))
            elif tag in ("ul", "ol"):
                for li in child.children:
                    if isinstance(li, _Node) and li.tag == "li":
                        writer().add_li(_inline_runs(li))
            elif tag == "table":
                title = state["writer"].base_title if state["writer"] else section_label
                _table_slide(prs, f"{title} — tabla", section_label, child)
            elif tag in ("p", "div", "span"):
                if _has_block_children(child):
                    walk(child)
                else:
                    runs = _inline_runs(child)
                    if runs:
                        writer().add_p(runs)
            else:
                walk(child)

    walk(root)


def render_pptx(sections: list[dict]) -> bytes:
    """`sections` is [{"label": <tab display name>, "html": <rendered
    innerHTML for that tab>}, ...] — same payload shape the "Descargar
    PDF" button already sends to help_report.render_pdf()."""
    from io import BytesIO
    from datetime import datetime
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _title_slide(prs, datetime.now().strftime("%d/%m/%Y %H:%M"))

    for sec in sections or []:
        label = (sec.get("label") or sec.get("tab") or "").strip()
        if not label:
            continue
        _divider_slide(prs, label)
        root = _parse_html(sec.get("html") or "")
        _render_section(prs, root, label)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
